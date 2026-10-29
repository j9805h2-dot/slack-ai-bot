import os
import re
import csv
import io
import json
import time
import base64
from collections import OrderedDict
import requests
from fastapi import FastAPI, Request, BackgroundTasks
from fastapi.responses import JSONResponse, Response
from slack_sdk import WebClient
from dotenv import load_dotenv

# 로컬 텍스트 추출 라이브러리
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

load_dotenv()

app = FastAPI()

SLACK_BOT_TOKEN = os.environ.get("SLACK_BOT_TOKEN")
slack_client = WebClient(token=SLACK_BOT_TOKEN) if SLACK_BOT_TOKEN else None

GEMINI_API_KEYS = []
for k, v in os.environ.items():
    if k.startswith("GEMINI_API_KEY") and v.strip():
        clean_key = v.strip().strip('"').strip("'")
        GEMINI_API_KEYS.append(clean_key)

api_key_index = 0
print(f"🔑 로드된 Gemini API 키 개수: {len(GEMINI_API_KEYS)}개")

MASTER_SHEET_ID = os.environ.get("MASTER_SHEET_ID", "1k-eGv2EFu4GG3H3oKG-Ql9IF-UVD_kZMTjnUiyLqLp4")
SHEET_CSV_URL = f"https://docs.google.com/spreadsheets/d/{MASTER_SHEET_ID}/export?format=csv"
DAESEONG_QNA_WEBAPP_URL = os.environ.get("DAESEONG_WEBAPP_URL") or os.environ.get("LOG_WEBHOOK_URL", "")

KNOWLEDGE_CACHE = []
CORRECTED_CACHE = []
LAST_SYNC_TIME = 0

def fetch_master_knowledge():
    global KNOWLEDGE_CACHE
    try:
        res = requests.get(SHEET_CSV_URL, timeout=10)
        res.raise_for_status()
        res.encoding = 'utf-8'
        rows = list(csv.reader(io.StringIO(res.text)))
        if not rows:
            return
        header = rows[0]
        KNOWLEDGE_CACHE = [
            {header[i].strip(): row[i].strip() for i in range(min(len(header), len(row)))}
            for row in rows[1:] if any(row)
        ]
        print(f"✅ [1/2] 📚AI_마스터지식 로드 완료: {len(KNOWLEDGE_CACHE)}건")
    except Exception as e:
        print(f"❌ 📚AI_마스터지식 로드 실패: {e}")

def fetch_corrected_knowledge():
    global CORRECTED_CACHE
    if not DAESEONG_QNA_WEBAPP_URL:
        return
    try:
        res = requests.get(DAESEONG_QNA_WEBAPP_URL, timeout=12)
        if res.status_code == 200:
            data = res.json()
            CORRECTED_CACHE = data.get("data", [])
            print(f"✅ [2/2] 💢대성_질의 실시간 관리자 보정 지식 로드 완료: {len(CORRECTED_CACHE)}건")
    except Exception as e:
        print(f"⚠️ 💢대성_질의 실시간 보정 지식 통신 실패: {e}")

def refresh_all_knowledge():
    global LAST_SYNC_TIME
    print("\n🔄 지식 베이스 실시간 갱신 시작...")
    fetch_master_knowledge()
    fetch_corrected_knowledge()
    LAST_SYNC_TIME = time.time()

refresh_all_knowledge()

def normalize_text(text: str) -> str:
    return re.sub(r'[\s\-_/&]+', '', text.lower())

def search_relevant_knowledge(query: str, top_k: int = 5) -> dict:
    global KNOWLEDGE_CACHE, CORRECTED_CACHE, LAST_SYNC_TIME
    if time.time() - LAST_SYNC_TIME > 60:
        refresh_all_knowledge()

    cleaned_query = normalize_text(query)
    tokens = [t for t in re.split(r'\s+', query) if len(t) >= 2]

    def score_items(items):
        results = []
        for item in items:
            score = 0
            item_str = " ".join([str(v) for v in item.values()])
            norm_str = normalize_text(item_str)
            if cleaned_query in norm_str:
                score += 10
            for t in tokens:
                if normalize_text(t) in norm_str:
                    score += 3
            if score > 0:
                results.append((score, item))
        results.sort(key=lambda x: x[0], reverse=True)
        return [it for _, it in results[:top_k]]

    return {
        "corrected_rules": score_items(CORRECTED_CACHE),
        "master_rules": score_items(KNOWLEDGE_CACHE)
    }

def convert_to_slack_markdown(text: str) -> str:
    return re.sub(r'\*\*(.*?)\*\*', r'*\1*', text)

def call_gemini_api(prompt: str) -> str:
    global api_key_index
    if not GEMINI_API_KEYS:
        return "⚠️ 등록된 Gemini API 키가 없습니다."

    headers = {"Content-Type": "application/json"}
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.2, "maxOutputTokens": 2048}
    }

    active_models = ["gemini-flash-latest", "gemini-2.0-flash"]
    max_attempts = len(GEMINI_API_KEYS) * len(active_models)
    last_error_msg = ""

    for _ in range(max_attempts):
        current_key = GEMINI_API_KEYS[api_key_index]
        for model in active_models:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={current_key}"
            try:
                res = requests.post(url, headers=headers, json=payload, timeout=25)
                if res.status_code == 200:
                    data = res.json()
                    raw_text = data["candidates"][0]["content"]["parts"][0]["text"]
                    return convert_to_slack_markdown(raw_text)
                elif res.status_code == 429:
                    api_key_index = (api_key_index + 1) % len(GEMINI_API_KEYS)
                    time.sleep(1)
                    break
                else:
                    last_error_msg = f"{model} ({res.status_code})"
            except Exception as e:
                last_error_msg = str(e)

        api_key_index = (api_key_index + 1) % len(GEMINI_API_KEYS)

    return f"죄송합니다. AI 서비스 연결에 일시적인 지연이 발생했습니다.\n(원인: {last_error_msg[:100]})"

# =========================================================================
# 💡 [핵심] 100% 로컬 텍스트 추출 (Gemini API 미호출 -> 429 원천 차단)
# =========================================================================
def extract_text_locally(file_bytes: bytes, file_name: str) -> str:
    lower = file_name.lower()
    text = ""
    try:
        if lower.endswith(".pdf"):
            if PdfReader is None:
                print("❌ pypdf 패키지가 설치되지 않았습니다. 설치를 진행해 주세요.")
                return ""
            reader = PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages):
                t = page.extract_text()
                if t:
                    text += f"\n[페이지 {i+1}]\n" + t.strip() + "\n"

        elif lower.endswith(".pptx"):
            if Presentation is None:
                print("❌ python-pptx 패키지가 설치되지 않았습니다. 설치를 진행해 주세요.")
                return ""
            prs = Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_text.append(paragraph.text.strip())
                if slide_text:
                    text += f"\n[슬라이드 {i+1}]\n" + "\n".join(slide_text) + "\n"
    except Exception as e:
        print(f"❌ 파일 파싱 예외 ({file_name}): {e}")

    return text.strip()

def bg_process_drive_sync(payload: dict):
    file_name = payload.get("fileName", "")
    b64_contents = payload.get("fileBase64", "")
    
    if not b64_contents:
        return

    file_bytes = base64.b64decode(b64_contents)
    print(f"🚀 [로컬 파싱 시작] {file_name} ({len(file_bytes)} bytes)")

    # Gemini API를 전혀 타지 않고 로컬에서 0.1초 만에 텍스트 추출
    extracted_text = extract_text_locally(file_bytes, file_name)

    if extracted_text and DAESEONG_QNA_WEBAPP_URL:
        save_payload = {
            "action": "append_master_knowledge",
            "title": f"[{file_name}] 업무 규정",
            "content": extracted_text,
            "source": f"드라이브문서({file_name})"
        }
        try:
            res = requests.post(DAESEONG_QNA_WEBAPP_URL, json=save_payload, timeout=25)
            print(f"✅ [마스터 시트 적재 완료] {file_name} ({len(extracted_text)}자) -> 상태: {res.status_code}")
        except Exception as e:
            print(f"❌ 마스터 시트 콜백 실패 ({file_name}): {e}")
    else:
        print(f"⚠️ [파싱 실패 또는 내용 없음] {file_name}")

@app.post("/sync/drive-file")
async def sync_drive_file(request: Request, background_tasks: BackgroundTasks):
    data = await request.json()
    background_tasks.add_task(bg_process_drive_sync, data)
    return {"status": "accepted", "message": f"{data.get('fileName')} 작업 등록 완료"}

# =========================================================================
# 슬랙 이벤트 & 인터랙션 처리 핸들러
# =========================================================================
def get_thread_context_and_count(channel: str, thread_ts: str) -> tuple:
    if not slack_client or not thread_ts:
        return "", "", 1
    try:
        res = slack_client.conversations_replies(channel=channel, ts=thread_ts, limit=20)
        messages = res.get("messages", [])
        if not messages:
            return "", "", 1

        first_q = messages[0].get("text", "")
        history = []
        user_question_count = 0

        for m in messages[:-1]:
            is_bot = "bot_id" in m
            if not is_bot:
                user_question_count += 1
            role = "AI" if is_bot else "상담원"
            history.append(f"{role}: {m.get('text', '')}")

        total_user_count = user_question_count + 1
        return first_q, "\n".join(history), total_user_count
    except Exception as e:
        print(f"⚠️ 스레드 문맥 조회 예외: {e}")
        return "", "", 1

def check_casual_intent(text: str) -> bool:
    clean = re.sub(r'[^가-힣a-zA-Z0-9]', '', text.lower())
    business_keywords = [
        "본인확인", "확인해야", "확인방법", "연체", "해약", "환급", "양도", 
        "명의", "프로세스", "절차", "서류", "발급", "신청", "인증", "기준", "어떻게"
    ]
    if any(b in clean for b in business_keywords):
        return False

    casual_keywords = [
        "안녕", "하이", "헬로", "hello", "hi", "좋은아침", "굿모닝", "퇴근", "출근",
        "고마워", "감사", "땡큐", "thanks", "수고", "고생", "화이팅", "파이팅", "힘내",
        "대박", "최고", "짱", "짱이다", "굿잡", "든든",
        "알겠", "오케이", "ok", "네", "넵", "넹", "응", "어", "ㅇㅇ", "굿", "굳", "좋아"
    ]
    if len(clean) <= 15 and any(k in clean for k in casual_keywords):
        return True
    return False

def check_manual_escalation_intent(text: str) -> bool:
    keywords = ["팀장", "관리자", "sv", "에스브이", "상급자", "물어볼", "물어봐야", "물어볼래", "확인해볼", "올려볼", "질의서", "질문있어", "질문 있어"]
    return any(k in text.lower() for k in keywords)

def log_to_google_sheet(gubun: str, thread_ts: str, user_id: str, first_q: str, sub_q: str, answer: str):
    if not DAESEONG_QNA_WEBAPP_URL:
        return
    try:
        requests.post(DAESEONG_QNA_WEBAPP_URL, json={
            "action": "log",
            "gubun": gubun,
            "thread_ts": thread_ts,
            "user_id": user_id,
            "first_question": first_q,
            "sub_question": sub_q,
            "answer": answer,
            "feedback": "",
            "review_status": "검수 필요"
        }, timeout=5)
    except Exception as e:
        print(f"⚠️ 질의로그 시트 전송 실패: {e}")

def build_modal_view_dict(channel_id: str, thread_ts: str, user_name: str, user_id: str) -> dict:
    return {
        "type": "modal",
        "callback_id": "submit_inquiry_modal",
        "title": {"type": "plain_text", "text": "질의 접수"},
        "submit": {"type": "plain_text", "text": "전송"},
        "close": {"type": "plain_text", "text": "취소"},
        "private_metadata": json.dumps({
            "channel_id": channel_id,
            "thread_ts": thread_ts,
            "user_id": user_id,
            "user_name": user_name
        }),
        "blocks": [
            {
                "type": "input",
                "block_id": "team_block",
                "label": {"type": "plain_text", "text": "소속팀 *"},
                "element": {"type": "plain_text_input", "action_id": "team_input", "placeholder": {"type": "plain_text", "text": "예: 1팀, 6팀"}}
            },
            {
                "type": "input",
                "block_id": "author_block",
                "label": {"type": "plain_text", "text": "상담사명(작성자) *"},
                "element": {"type": "plain_text_input", "action_id": "author_input", "initial_value": user_name or "", "placeholder": {"type": "plain_text", "text": "성함을 입력하세요."}}
            },
            {
                "type": "input",
                "block_id": "contract_block",
                "label": {"type": "plain_text", "text": "계약번호 / 연락처 *"},
                "element": {"type": "plain_text_input", "action_id": "contract_input", "placeholder": {"type": "plain_text", "text": "계약번호 또는 전화번호"}}
            },
            {
                "type": "input",
                "block_id": "name_block",
                "label": {"type": "plain_text", "text": "회원명 *"},
                "element": {"type": "plain_text_input", "action_id": "name_input", "placeholder": {"type": "plain_text", "text": "회원명을 입력하세요."}}
            },
            {
                "type": "input",
                "block_id": "customerQ_block",
                "label": {"type": "plain_text", "text": "고객 문의 내용(고객) *"},
                "element": {"type": "plain_text_input", "action_id": "customerQ_input", "placeholder": {"type": "plain_text", "text": "고객이 인입 시 문의한 핵심 내용"}}
            },
            {
                "type": "input",
                "block_id": "myQ_block",
                "label": {"type": "plain_text", "text": "문의 내용(내가 궁금한 점) *"},
                "element": {"type": "plain_text_input", "action_id": "myQ_input", "multiline": True, "placeholder": {"type": "plain_text", "text": "확인받고자 하는 구체적인 내용 작성"}}
            }
        ]
    }

def process_slack_question(channel: str, ts: str, thread_ts: str, user_text: str, user_id: str):
    is_thread_reply = bool(thread_ts and thread_ts != ts)
    target_ts = thread_ts if is_thread_reply else ts

    if slack_client:
        try:
            slack_client.reactions_add(channel=channel, timestamp=ts, name="eyes")
        except Exception:
            pass

    clean_prompt = re.sub(r'^\[질의-AI\]\s*', '', user_text).strip()

    if check_casual_intent(clean_prompt):
        casual_prompt = f"상담원의 인사나 반응에 다정하게 화답해 주세요. 상담원 말: '{clean_prompt}'"
        answer = call_gemini_api(casual_prompt)
        if slack_client:
            try:
                slack_client.reactions_remove(channel=channel, timestamp=ts, name="eyes")
                slack_client.chat_postMessage(channel=channel, thread_ts=target_ts, text=answer)
                slack_client.reactions_add(channel=channel, timestamp=ts, name="white_check_mark")
            except Exception:
                pass
        log_to_google_sheet("일상대화", target_ts, user_id, clean_prompt, "", answer)
        return

    first_q_text, thread_history, user_q_count = ("", "", 1)
    if is_thread_reply:
        first_q_text, thread_history, user_q_count = get_thread_context_and_count(channel, thread_ts)
        first_q_clean = re.sub(r'^\[질의-AI\]\s*', '', first_q_text).strip()
    else:
        first_q_clean = clean_prompt

    wants_manual_escalate = check_manual_escalation_intent(clean_prompt)
    is_escalation_triggered = (user_q_count >= 3) or wants_manual_escalate

    if wants_manual_escalate:
        answer = "관리자(SV)님께 바로 질문하실 수 있도록 접수를 도와드리겠습니다.\n아래 *[📋 관리자(SV) 질의 접수]* 버튼을 눌러 질의서를 작성해 주세요!"
    else:
        search_target = f"{first_q_clean} {clean_prompt}" if is_thread_reply else clean_prompt
        knowledge = search_relevant_knowledge(search_target)

        corrected_context = ""
        if knowledge.get('corrected_rules'):
            corrected_context = f"[⚠️ 관리자 최우선 보정 지침]:\n{json.dumps(knowledge['corrected_rules'], ensure_ascii=False, indent=2)}\n"

        system_instruction = f"""
당신은 보람상조 양평센터 상담원들을 지원하는 사내 AI 어시스턴트입니다.
사내 규정 지식과 관리자 보정 지침을 바탕으로 실무적이고 정확하게 답변하세요.

{corrected_context}

[기본 사내 규정 지식]:
{json.dumps(knowledge['master_rules'], ensure_ascii=False, indent=2)}

[이전 대화 맥락]:
{thread_history if thread_history else "신규 질문입니다."}

[질문]:
{clean_prompt}

[원칙]:
1. 관리자 보정 지침이 있으면 1순위로 적용하세요.
2. 질문에 '본인확인'이 포함되면 본인확인 필수 항목(성함, 생년월일, 연락처 등)을 명시하세요.
3. '절차', '순서', '방법' 질문은 1단계 -> 2단계 순서로 먼저 명확히 안내하세요.
"""
        base_answer = call_gemini_api(system_instruction)
        answer = f"{base_answer}\n\n---\n*안내*: 세부 확인이 필요하시면 아래 버튼으로 관리자 질의를 접수하세요." if user_q_count >= 3 else base_answer

    if slack_client:
        try:
            slack_client.reactions_remove(channel=channel, timestamp=ts, name="eyes")
        except Exception:
            pass

        if is_escalation_triggered:
            blocks = [
                {"type": "section", "text": {"type": "mrkdwn", "text": answer}},
                {"type": "divider"},
                {
                    "type": "actions",
                    "elements": [
                        {"type": "button", "text": {"type": "plain_text", "text": "📋 관리자(SV) 질의 접수"}, "style": "primary", "action_id": "open_inquiry_modal_btn", "value": target_ts},
                        {"type": "button", "text": {"type": "plain_text", "text": "❌ 작성 취소 (해결됨)"}, "action_id": "dismiss_inquiry_btn", "value": target_ts}
                    ]
                }
            ]
            post_kwargs = {"blocks": blocks}
        else:
            post_kwargs = {}

        try:
            slack_client.chat_postMessage(channel=channel, thread_ts=target_ts, text=answer, **post_kwargs)
            slack_client.reactions_add(channel=channel, timestamp=ts, name="white_check_mark")
        except Exception as e:
            print(f"❌ 슬랙 발송 실패: {e}")

    gubun_val = f"대댓글({user_q_count}회차)" if is_thread_reply else "최초답변"
    sub_q_val = clean_prompt if is_thread_reply else ""
    log_to_google_sheet(gubun_val, target_ts, user_id, first_q_clean, sub_q_val, answer)

PROCESSED_EVENT_IDS = OrderedDict()

@app.post("/slack/events")
async def slack_events(request: Request, background_tasks: BackgroundTasks):
    body = await request.body()
    data = json.loads(body.decode("utf-8"))

    if "challenge" in data:
        return {"challenge": data["challenge"]}

    event_id = data.get("event_id")
    if event_id:
        if event_id in PROCESSED_EVENT_IDS:
            return {"status": "already_processed"}
        PROCESSED_EVENT_IDS[event_id] = True
        if len(PROCESSED_EVENT_IDS) > 1000:
            PROCESSED_EVENT_IDS.popitem(last=False)

    event = data.get("event", {})
    text = event.get("text", "")
    channel = event.get("channel")
    ts = event.get("ts")
    thread_ts = event.get("thread_ts")
    bot_id = event.get("bot_id")
    user_id = event.get("user", "상담원")

    if not bot_id and (text.startswith("[질의-AI]") or (thread_ts and thread_ts != ts)):
        background_tasks.add_task(process_slack_question, channel, ts, thread_ts, text, user_id)

    return {"status": "ok"}

def bg_process_submission(view: dict):
    values = view.get("state", {}).get("values", {})
    meta = json.loads(view.get("private_metadata", "{}"))

    sheet_payload = {
        "action": "escalate",
        "team": values.get("team_block", {}).get("team_input", {}).get("value", ""),
        "author": f"<@{meta.get('user_id', '')}>" if meta.get("user_id") else values.get("author_block", {}).get("author_input", {}).get("value", ""),
        "contract": values.get("contract_block", {}).get("contract_input", {}).get("value", ""),
        "name": values.get("name_block", {}).get("name_input", {}).get("value", ""),
        "customerQ": values.get("customerQ_block", {}).get("customerQ_input", {}).get("value", ""),
        "myQ": values.get("myQ_block", {}).get("myQ_input", {}).get("value", ""),
        "slack_thread_key": f"{meta.get('channel_id')}|{meta.get('thread_ts')}"
    }

    if DAESEONG_QNA_WEBAPP_URL:
        try:
            requests.post(DAESEONG_QNA_WEBAPP_URL, json=sheet_payload, timeout=5)
        except Exception as e:
            print(f"⚠️ 대장 전송 실패: {e}")

    channel_id = meta.get("channel_id")
    thread_ts = meta.get("thread_ts")
    if slack_client and channel_id and thread_ts:
        try:
            slack_client.chat_postMessage(channel=channel_id, thread_ts=thread_ts, text="전송 완료. 관리자가 확인하고 답변을 드릴 때 까지 1~2일 정도 소요될 수 있는 점 양해 부탁 드려요~")
        except Exception:
            pass

@app.post("/slack/interactive")
async def slack_interactive(request: Request, background_tasks: BackgroundTasks):
    try:
        form_data = await request.form()
        payload_raw = form_data.get("payload")
        if not payload_raw:
            return Response(status_code=200)

        payload = json.loads(payload_raw)

        if payload.get("type") == "block_actions":
            for action in payload.get("actions", []):
                action_id = action.get("action_id")
                raw_val = action.get("value", "")

                if action_id == "open_inquiry_modal_btn":
                    trigger_id = payload.get("trigger_id")
                    channel_id = payload.get("channel", {}).get("id")
                    user_info = payload.get("user", {})
                    user_id = user_info.get("id")
                    user_name = user_info.get("name", "상담사")

                    if slack_client and trigger_id:
                        modal_view = build_modal_view_dict(channel_id, raw_val, user_name, user_id)
                        slack_client.views_open(trigger_id=trigger_id, view=modal_view)
                    return Response(status_code=200)

                elif action_id == "dismiss_inquiry_btn":
                    channel_id = payload.get("channel", {}).get("id")
                    user_id = payload.get("user", {}).get("id")
                    msg_ts = payload.get("message", {}).get("ts")
                    orig_text = payload.get("message", {}).get("text", "")

                    if slack_client and channel_id:
                        if msg_ts:
                            slack_client.chat_update(channel=channel_id, ts=msg_ts, text=orig_text, blocks=[{"type": "section", "text": {"type": "mrkdwn", "text": orig_text}}])
                        slack_client.chat_postMessage(channel=channel_id, thread_ts=raw_val, text=f"👍 <@{user_id}> 상담원님, 문의가 해결되어 질의서 접수를 종료합니다.")
                    return Response(status_code=200)

        if payload.get("type") == "view_submission":
            view = payload.get("view", {})
            if view.get("callback_id") == "submit_inquiry_modal":
                background_tasks.add_task(bg_process_submission, view)
                return JSONResponse(content={"response_action": "clear"})

    except Exception as err:
        print(f"❌ interactive 예외: {err}")
        return Response(status_code=200)

    return Response(status_code=200)
